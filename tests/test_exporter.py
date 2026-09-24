"""Testes da camada de exportação desacoplada."""

from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

import rex.exporter as exporter_module
from rex.errors import ExportError
from rex.exporter import export_result
from rex.models import ExtractionResult
from rex.workbook import WORKBOOK_MODE_PER_SAMPLE


def test_exports_structured_result_without_engine_dependency(tmp_path: Path):
    result = ExtractionResult(
        data=pd.DataFrame(
            [
                {
                    "Amostra": "A01 externo 071815",
                    "Micrografia": "media/image1.png",
                    "Serie": 1,
                    "Ponto": 1,
                    "Elemento": "Si",
                    "TipoLinha": "K series",
                    "PercentualPeso": 59.57,
                    "SigmaPeso": 0.23,
                    "PercentualAtomico": 73.86,
                }
            ]
        ),
        samples=pd.DataFrame(
            [
                {
                    "SampleIndex": 1,
                    "SampleName": "A01 externo 071815",
                    "ImageFile": "media/image1.png",
                }
            ]
        ),
    )

    paths = export_result(result, tmp_path, source_path="Laudo A.docx")

    assert paths.excel.is_file()
    assert paths.excel.name.startswith("Laudo A_REX-")
    assert paths.excel.name.endswith(".xlsx")
    assert paths.csv.is_file()
    assert paths.csv.name == "dados_extraidos.csv"


def test_csv_stays_consolidated_with_per_sample_workbook(tmp_path: Path):
    result = ExtractionResult(
        data=pd.DataFrame(
            {"Amostra": ["A01", "A02"], "Elemento": ["Si", "Al"]}
        ),
        samples=pd.DataFrame({"SampleName": ["A01", "A02"]}),
    )

    paths = export_result(
        result,
        tmp_path,
        source_path="laudo.docx",
        workbook_mode=WORKBOOK_MODE_PER_SAMPLE,
    )

    assert pd.read_csv(paths.csv)["Amostra"].tolist() == ["A01", "A02"]
    workbook = load_workbook(paths.excel, read_only=True)
    try:
        assert workbook.sheetnames == ["AmostrasDetectadas", "A01", "A02"]
    finally:
        workbook.close()


def test_failed_export_preserves_previous_files(tmp_path: Path, monkeypatch):
    previous_excel = tmp_path / "laudo_REX-20260923-120000-000000.xlsx"
    previous_csv = tmp_path / "dados_extraidos.csv"
    previous_excel.write_bytes(b"planilha-anterior")
    previous_csv.write_text("csv-anterior", encoding="utf-8")
    result = ExtractionResult(data=pd.DataFrame(), samples=pd.DataFrame())

    def fail_workbook(*args, **kwargs):
        raise OSError("disco indisponível")

    monkeypatch.setattr(exporter_module, "write_excel_workbook", fail_workbook)
    monkeypatch.setattr(exporter_module, "datetime", _fixed_datetime)

    with pytest.raises(ExportError, match="Falha ao exportar"):
        export_result(result, tmp_path, source_path="laudo.docx")

    assert previous_excel.read_bytes() == b"planilha-anterior"
    assert previous_csv.read_text(encoding="utf-8") == "csv-anterior"
    assert not list(tmp_path.glob(".rex-export-*"))


def test_second_replacement_failure_restores_previous_pair(tmp_path: Path, monkeypatch):
    previous_excel = tmp_path / "laudo_REX-20260923-120000-000000.xlsx"
    previous_csv = tmp_path / "dados_extraidos.csv"
    previous_excel.write_bytes(b"planilha-anterior")
    previous_csv.write_text("csv-anterior", encoding="utf-8")
    result = ExtractionResult(data=pd.DataFrame({"Amostra": ["A01"]}), samples=pd.DataFrame())
    real_replace = exporter_module.os.replace

    def fail_csv_replacement(source, destination):
        if Path(destination) == previous_csv:
            raise OSError("falha na substituição do CSV")
        return real_replace(source, destination)

    monkeypatch.setattr(exporter_module.os, "replace", fail_csv_replacement)
    monkeypatch.setattr(exporter_module, "datetime", _fixed_datetime)

    with pytest.raises(ExportError, match="falha na substituição do CSV"):
        export_result(result, tmp_path, source_path="laudo.docx")

    assert previous_excel.read_bytes() == b"planilha-anterior"
    assert previous_csv.read_text(encoding="utf-8") == "csv-anterior"


class _fixed_datetime:
    @staticmethod
    def now():
        from datetime import datetime

        return datetime(2026, 9, 23, 12, 0, 0)
