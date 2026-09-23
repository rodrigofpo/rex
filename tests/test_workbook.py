"""Testes da estrutura das pastas de trabalho Excel."""

import zipfile
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from rex.workbook import (
    WORKBOOK_MODE_CONSOLIDATED,
    WORKBOOK_MODE_PER_SAMPLE,
    unique_sheet_name,
    validate_workbook_mode,
    write_excel_workbook,
)


def _sample_dataframes() -> tuple[pd.DataFrame, pd.DataFrame]:
    data = pd.DataFrame(
        [
            {
                "Amostra": "A01 externo 071815",
                "Micrografia": "media/image1.png",
                "Serie": 1,
                "Ponto": 1,
                "Elemento": "Si",
                "TipoLinha": "K series",
                "PercentualPeso": 59.57,
                "SigmaPeso": 0.23,
                "PercentualAtomico": 73.86,
            },
            {
                "Amostra": "A02 interno 071816",
                "Micrografia": "media/image2.png",
                "Serie": 1,
                "Ponto": 2,
                "Elemento": "Al",
                "TipoLinha": "K series",
                "PercentualPeso": 40.43,
                "SigmaPeso": 0.19,
                "PercentualAtomico": 26.14,
            },
        ]
    )
    summary = pd.DataFrame(
        [
            {"SampleIndex": 1, "SampleName": "A01 externo 071815", "ImageFile": "media/image1.png"},
            {"SampleIndex": 2, "SampleName": "A02 interno 071816", "ImageFile": "media/image2.png"},
        ]
    )
    return data, summary


def test_writes_consolidated_workbook(tmp_path: Path):
    data, summary = _sample_dataframes()
    output = write_excel_workbook(
        data,
        summary,
        tmp_path / "consolidated.xlsx",
        WORKBOOK_MODE_CONSOLIDATED,
    )

    workbook = load_workbook(output, read_only=True, data_only=True)
    try:
        assert workbook.sheetnames == ["DadosCompletos", "AmostrasDetectadas"]
    finally:
        workbook.close()


def test_writes_one_sheet_per_sample(tmp_path: Path):
    data, summary = _sample_dataframes()
    output = write_excel_workbook(
        data,
        summary,
        tmp_path / "nested" / "per-sample.xlsx",
        WORKBOOK_MODE_PER_SAMPLE,
    )

    workbook = load_workbook(output, read_only=True, data_only=True)
    try:
        assert workbook.sheetnames == [
            "AmostrasDetectadas",
            "A01 externo 071815",
            "A02 interno 071816",
        ]
        assert workbook["A01 externo 071815"].max_row == 2
        assert workbook["A02 interno 071816"].max_row == 2
        first_sample_row = list(
            workbook["A01 externo 071815"].iter_rows(min_row=2, values_only=True)
        )[0]
        assert first_sample_row[0] == "A01 externo 071815"
        assert first_sample_row[3] == 1
    finally:
        workbook.close()

    with zipfile.ZipFile(output) as archive:
        assert archive.testzip() is None


def test_sheet_names_are_valid_limited_and_unique():
    used_names: set[str] = set()

    first = unique_sheet_name("Amostra/01:*?[] com nome muito longo 071815", used_names)
    second = unique_sheet_name("Amostra\\01:*?[] com nome muito longo 071815", used_names)

    assert len(first) <= 31
    assert len(second) <= 31
    assert first != second
    assert not set("\\/*?:[]").intersection(first)


def test_rejects_unknown_workbook_mode():
    try:
        validate_workbook_mode("unknown")
    except ValueError as exc:
        assert "Modo de pasta de trabalho inválido" in str(exc)
    else:
        raise AssertionError("O modo inválido deveria ser rejeitado.")
